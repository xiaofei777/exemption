from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
OUT='/Users/cfy/code/推免/计算机保研408专业课速成题库.docx'
doc=Document(); sec=doc.sections[0]; sec.top_margin=Inches(.7); sec.bottom_margin=Inches(.7); sec.left_margin=Inches(.8); sec.right_margin=Inches(.8)
for n,size,col in [('Normal',10.5,'222222'),('Title',26,'17365D'),('Heading 1',18,'1F4E79'),('Heading 2',13,'2F75B5')]:
 s=doc.styles[n]; s.font.name='Songti SC'; s._element.rPr.rFonts.set(qn('w:ascii'),'Songti SC'); s._element.rPr.rFonts.set(qn('w:hAnsi'),'Songti SC'); s._element.rPr.rFonts.set(qn('w:eastAsia'),'宋体'); s.font.size=Pt(size); s.font.color.rgb=RGBColor.from_string(col); s.font.bold=n!='Normal'; s.paragraph_format.space_after=Pt(4)
def h(t,l=1): doc.add_paragraph(t,style=f'Heading {l}')
def q(a,b):
 p=doc.add_paragraph(); r=p.add_run('Q  '+a); r.bold=True; r.font.color.rgb=RGBColor(31,78,121)
 p=doc.add_paragraph('答  '+b); p.paragraph_format.left_indent=Inches(.18); p.paragraph_format.space_after=Pt(4)
def bullets(items):
 for t in items:
  p=doc.add_paragraph(style='List Bullet'); p.paragraph_format.left_indent=Inches(.22); p.paragraph_format.first_line_indent=Inches(-.12); p.add_run(t)
title=doc.add_paragraph(); title.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=title.add_run('计算机保研推免面试\n408 专业课速成题库'); r.bold=True; r.font.size=Pt(26); r.font.color.rgb=RGBColor(23,54,93)
p=doc.add_paragraph('面向面试表达的高频问题 · 先背结论，再补原理'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
h('使用说明',1); bullets(['复习顺序：数据结构 → 操作系统 → 计算机网络 → 组成原理。','回答结构：定义 / 核心区别 → 原理 → 应用场景或优缺点。','时间有限时，优先背每节 Q 题；每题准备 30 秒版本和 2 分钟版本。'])
h('一、数据结构（优先级最高）',1); h('排序算法',2); bullets(['冒泡 / 插入：O(n²)，稳定；插入排序适合基本有序数据。','选择：O(n²)，不稳定，交换次数少。','快速：平均 O(nlogn)，最坏 O(n²)，不稳定；可随机选主元优化。','归并：O(nlogn)，稳定，但需要 O(n) 额外空间。','堆：O(nlogn)，不稳定，额外空间 O(1)。']); q('快速排序为什么最坏是 O(n²)？','每次选到最大或最小元素，划分极不平衡。'); q('哪些排序算法稳定？','冒泡、插入、归并；快速、选择、堆通常不稳定。')
h('树、图与查找',2); q('B 树和 B+ 树有什么区别？','B+ 树的数据都在叶子节点，叶子节点通过指针连接，范围查询和磁盘访问更高效。'); q('AVL 树和红黑树有什么区别？','AVL 平衡更严格，查询快但旋转多；红黑树插入删除更高效。'); q('DFS 和 BFS 的区别？','DFS 用栈或递归，适合回溯；BFS 用队列，适合无权图最短路径。'); q('Dijkstra 能处理负权边吗？','不能，要求边权非负。'); q('哈希冲突如何解决？','链地址法、开放寻址法、再哈希法。'); q('栈和队列的典型应用？','栈：括号匹配、函数调用、DFS；队列：BFS、任务调度、缓冲区。')
h('二、操作系统',1); q('进程和线程的区别？','进程是资源分配基本单位，线程是 CPU 调度基本单位；同进程线程共享地址空间，但有独立栈和寄存器。'); q('死锁四个必要条件？','互斥、占有并等待、不可剥夺、循环等待。'); q('死锁如何解决？','预防、避免、检测与解除；银行家算法属于避免。'); q('分页和分段区别？','分页固定大小，服务于内存管理；分段按逻辑划分、大小可变。'); q('什么是虚拟内存？','利用磁盘和部分内存提供更大的地址空间，通常通过分页实现。'); q('FIFO、LRU、OPT 有何区别？','FIFO 淘汰最早进入页面，可能 Belady 异常；LRU 淘汰最长未使用；OPT 淘汰未来最长不用页面，理论最优。'); q('select、poll、epoll 区别？','select/poll 需遍历文件描述符；epoll 通过事件通知，更适合大量连接。')
h('三、计算机网络',1); q('TCP 和 UDP 的区别？','TCP 面向连接、可靠、有序，具备流量与拥塞控制；UDP 无连接、开销小、延迟低但不保证可靠。'); q('TCP 三次握手为什么是三次？','确认双方收发能力并同步序列号，避免历史请求误建立。'); q('TCP 四次挥手为什么是四次？','全双工连接的两个方向需要分别关闭。'); q('TIME_WAIT 的作用？','确保最后 ACK 到达，并让旧连接延迟报文消失。'); q('TCP 如何保证可靠传输？','序列号、确认应答、超时重传、校验和、滑动窗口。'); q('流量控制和拥塞控制的区别？','流量控制针对接收方能力；拥塞控制针对网络整体状况。'); q('HTTP 和 HTTPS 的区别？','HTTPS=HTTP+TLS，提供加密、认证和完整性保护。'); q('常见 HTTP 状态码？','200 成功；301/302 重定向；400 请求错误；401 未认证；403 无权限；404 不存在；500 服务器错误。')
h('四、计算机组成原理',1); q('Cache 为什么有效？','程序具有时间和空间局部性，Cache 保存近期及附近数据，减少访问主存时间。'); q('Cache 映射方式？','直接映射、全相联映射、组相联映射。'); q('一条指令的执行过程？','取指、译码、执行、访存、写回。'); q('什么是流水线？','将指令划分为多个阶段，让多条指令重叠执行；有结构、数据、控制冲突。'); q('中断和异常区别？','中断通常来自外部硬件；异常通常由当前指令执行引起。'); q('中断和 DMA 区别？','中断需 CPU 参与传输；DMA 可直接在内存和外设间传输。'); q('SRAM 和 DRAM 区别？','SRAM 快、贵、不需刷新；DRAM 便宜、密度高、需刷新。'); q('为什么 GPU 适合深度学习？','深度学习含大量可并行矩阵运算，GPU 有大量并行单元和更高带宽。')
h('五、一天速成清单',1); bullets(['上午：排序、B+ 树、AVL/红黑树、DFS/BFS、Dijkstra。','下午：进程线程、死锁、虚拟内存、页面置换、同步互斥。','晚上：TCP/UDP、握手挥手、HTTP/HTTPS、Cache、流水线。','最后 30 分钟：用“定义 → 原理 → 场景”把每道题说一遍。'])
p=doc.add_paragraph(); r=p.add_run('面试提醒：不会时先说已知定义，再说明推理路径，不要直接沉默。'); r.bold=True; r.font.color.rgb=RGBColor(192,80,77)
doc.save(OUT); print(OUT)
